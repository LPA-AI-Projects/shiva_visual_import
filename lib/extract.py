#!/usr/bin/env python3
"""
extract.py <source.pptx> <out.json>

Extracts a structured, per-slide representation of a PPTX file: title,
body text runs (with bullet/indent level), tables, images (with
normalized position), and a simple multi-column layout hint based on
x-position clustering of text boxes.

This is the deterministic "read the source deck" half of the pipeline.
The output JSON is what gets handed to Claude for semantic mapping onto
the brand component library — Claude never sees the raw PPTX.
"""
import sys, json, os
from pptx import Presentation

def emu_to_in(v):
    return round(v / 914400, 3) if v is not None else None

def max_font_size(text_frame):
    """Largest explicit font size (pt) found among the runs, or 0 if none set."""
    best = 0
    for p in text_frame.paragraphs:
        for r in p.runs:
            try:
                if r.font.size is not None:
                    best = max(best, r.font.size.pt)
            except Exception:
                pass
    return best

def title_score(shape, slide_h_in):
    """Higher = more likely to be the slide's main title. Font size is the
    dominant signal (a small eyebrow/kicker label near the top would win a
    pure position-based heuristic, but titles are reliably the largest text
    near the top of the slide)."""
    if not getattr(shape, 'has_text_frame', False) or not shape.text_frame.text.strip():
        return -1
    text = shape.text_frame.text.strip()
    if len(text.split()) > 18:  # titles are short; long text is body content
        return -1
    try:
        y_in = emu_to_in(shape.top)
    except Exception:
        y_in = None
    if y_in is None or y_in > slide_h_in * 0.55:
        return -1
    size = max_font_size(shape.text_frame)
    position_bonus = max(0, 1.5 - y_in)
    return size * 2 + position_bonus

def extract_text_frame(tf):
    paras = []
    for p in tf.paragraphs:
        text = ''.join(r.text for r in p.runs) or (p.text or '')
        text = text.strip()
        if not text:
            continue
        paras.append({'text': text, 'level': p.level or 0})
    return paras

def extract_table(shape):
    rows = []
    for row in shape.table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return rows

def extract_slide(slide, idx, slide_h_in, img_out_dir):
    result = {
        'index': idx, 'title': None, 'text_blocks': [], 'tables': [],
        'images': [], 'shape_count': 0,
    }

    # Pass 1: find the best title candidate by score (font size dominant).
    # Compare by shape_id, not object identity — python-pptx creates a new
    # wrapper object on every iteration of slide.shapes, so `shape is
    # best_shape` across two separate loops silently never matches.
    best_shape_id, best_score = None, -1
    for shape in slide.shapes:
        sc = title_score(shape, slide_h_in)
        if sc > best_score:
            best_score, best_shape_id = sc, shape.shape_id

    for shape in slide.shapes:
        result['shape_count'] += 1
        try:
            x, y, w, h = emu_to_in(shape.left), emu_to_in(shape.top), emu_to_in(shape.width), emu_to_in(shape.height)
        except Exception:
            x = y = w = h = None

        if getattr(shape, 'has_table', False):
            try:
                rows = extract_table(shape)
                if any(any(c.strip() for c in r) for r in rows):
                    result['tables'].append({'x': x, 'y': y, 'w': w, 'h': h, 'rows': rows})
            except Exception:
                pass
            continue

        if shape.shape_type == 13:  # PICTURE
            try:
                img = shape.image
                fname = f"slide{idx}_img{result['shape_count']}.{img.ext}"
                with open(os.path.join(img_out_dir, fname), 'wb') as f:
                    f.write(img.blob)
                result['images'].append({'x': x, 'y': y, 'w': w, 'h': h, 'path': fname})
            except Exception:
                pass
            continue

        if getattr(shape, 'has_text_frame', False) and shape.text_frame.text.strip():
            paras = extract_text_frame(shape.text_frame)
            if not paras:
                continue
            if shape.shape_id == best_shape_id and result['title'] is None:
                result['title'] = ' '.join(p['text'] for p in paras)
            else:
                result['text_blocks'].append({'x': x, 'y': y, 'w': w, 'h': h, 'paragraphs': paras})
            continue

        if shape.shape_type == 6:  # GROUP — recurse one level deep
            try:
                for sub in shape.shapes:
                    if getattr(sub, 'has_text_frame', False) and sub.text_frame.text.strip():
                        paras = extract_text_frame(sub.text_frame)
                        if paras:
                            result['text_blocks'].append({'x': x, 'y': y, 'w': w, 'h': h, 'paragraphs': paras, 'from_group': True})
            except Exception:
                pass

    xs = sorted(set(round(b['x'], 1) for b in result['text_blocks'] if b['x'] is not None))
    result['column_hint'] = len(xs) if 1 < len(xs) <= 4 else 1
    return result

def main():
    src, out_json = sys.argv[1], sys.argv[2]
    img_out_dir = os.path.join(os.path.dirname(out_json), 'images')
    os.makedirs(img_out_dir, exist_ok=True)

    prs = Presentation(src)
    slide_h_in = emu_to_in(prs.slide_height)

    slides = [extract_slide(slide, i, slide_h_in, img_out_dir) for i, slide in enumerate(prs.slides, start=1)]

    out = {
        'source_slide_count': len(slides),
        'source_width_in': emu_to_in(prs.slide_width),
        'source_height_in': slide_h_in,
        'slides': slides,
    }
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    print(f"Extracted {len(slides)} slides -> {out_json}")

if __name__ == '__main__':
    main()
